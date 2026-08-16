"""One-off importer: reads 'Inserção_de dados_Mensal_ROV' from the source workbook
and upserts MonthlyEntry rows.

Run inside the container:
    docker cp /path/to/file.xlsx indicebbs_app:/tmp/file.xlsx
    docker exec -it indicebbs_app python import_xlsx.py /tmp/file.xlsx
"""
import sys
from openpyxl import load_workbook
from database import SessionLocal
from models import Rov, MonthlyEntry


SHEET = "Inserção_de dados_Mensal_ROV"


def run(path: str):
    wb = load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"ERROR: sheet {SHEET!r} not found. Available: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[SHEET]

    db = SessionLocal()
    try:
        rovs_by_id = {r.rov_id: r for r in db.query(Rov).all()}
        created = updated = skipped = 0

        for row in ws.iter_rows(values_only=True, min_row=2):
            year, month, rov_id, *rest = row + (None,) * 20
            if not year or not month or not rov_id:
                skipped += 1
                continue
            _vessel_auto, _contractor_auto, total_bbs, nc, _conf, _kpi, _pct_nc, c1, c2, c3, c4_crm, c4_elo, obs, *_ = rest

            rov = rovs_by_id.get(rov_id.strip() if isinstance(rov_id, str) else rov_id)
            if not rov:
                print(f"  WARN: ROV {rov_id!r} not in cadastro — skipped")
                skipped += 1
                continue

            total = int(total_bbs or 0)
            ncv = int(nc or 0)
            if ncv > total:
                ncv = total

            existing = (
                db.query(MonthlyEntry)
                .filter(
                    MonthlyEntry.year == int(year),
                    MonthlyEntry.month == int(month),
                    MonthlyEntry.rov_id == rov.id,
                )
                .first()
            )
            data = dict(
                total_bbs=total,
                non_conforming=ncv,
                c1=int(c1 or 0), c2=int(c2 or 0), c3=int(c3 or 0),
                c4_crm=int(c4_crm or 0), c4_elo=int(c4_elo or 0),
                observations=(obs.strip() if isinstance(obs, str) and obs.strip() else None),
            )
            if existing:
                for k, v in data.items():
                    setattr(existing, k, v)
                updated += 1
            else:
                db.add(MonthlyEntry(year=int(year), month=int(month), rov_id=rov.id, **data))
                created += 1

        db.commit()
        print(f"[import_xlsx] created={created} updated={updated} skipped={skipped}")
    finally:
        db.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python import_xlsx.py <path-to-file.xlsx>")
        sys.exit(2)
    run(sys.argv[1])
